"""文档解析：PDF / DOCX / MD / TXT → 纯文本（全程内存处理，不落盘临时文件）。"""
import io
from pathlib import Path


def extract_text(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _pdf(data)
    if name.endswith(".docx"):
        return _docx(data)
    if name.endswith((".md", ".markdown", ".txt")):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"不支持的文件类型: {filename}（支持 PDF/DOCX/MD/TXT）")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages)


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [para.text for para in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text_from_file(path: str | Path) -> str:
    """本地文件入口（调试用）。"""
    p = Path(path)
    return extract_text(p.read_bytes(), p.name)
